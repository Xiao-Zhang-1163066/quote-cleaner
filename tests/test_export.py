"""Tests for src/export.py.

The pure functions (build_comparison_rows, cheapest_supplier, supplier_grand_totals) get
normal unit tests, the same way rules.py's do -- build SupplierQuote/SupplierLine by hand,
no database. The openpyxl writing is covered by one integration test: write a real file to
tmp_path, then read it back with load_workbook and check the cells a human demo would look
at (headers, the highlighted "cheapest" cell, issue row colours). That test cares about
wiring, not pixels -- it would not catch a slightly-wrong shade of green.
"""
from decimal import Decimal as D
from pathlib import Path

from openpyxl import load_workbook

from src.export import (
    ComparisonRow, FILL_CHEAPEST, FILL_ERROR, FILL_WARNING, Item, SupplierPrice,
    build_comparison_rows, cheapest_supplier, export_workbook, group_missing_items,
    supplier_grand_totals,
)
from src.rules import Issue, Severity, SupplierLine, SupplierQuote


GIB = Item("GIB-STD-10", "10mm standard plasterboard", "sheet")
MESH = Item("MESH-665", "665 reinforcing mesh", "sheet")


def _quote(supplier: str, prices: dict[str, tuple[D, D]]) -> SupplierQuote:
    """Build a SupplierQuote with one line per (item_code -> (unit_price, line_total)) pair.
    Raw fields are irrelevant to export.py, so they're left None -- only rules.py's
    RAW-number checks need them.
    """
    lines = [
        SupplierLine(item_code=code, raw_qty=None, raw_unit_price=None, raw_line_total=None,
                    qty_canonical=None, unit_price_ex_gst=price, line_total_ex_gst=total)
        for code, (price, total) in prices.items()
    ]
    return SupplierQuote(supplier=supplier, gst_status="excl", stated_total=None, lines=lines)


# --- build_comparison_rows ---------------------------------------------------

def test_pivots_long_lines_into_one_row_per_item():
    harbour = _quote("Harbour", {"GIB-STD-10": (D("65.20"), D("652.00"))})
    kiwi = _quote("Kiwi Frame", {"GIB-STD-10": (D("68.00"), D("680.00"))})

    rows = build_comparison_rows([GIB], [harbour, kiwi])

    assert len(rows) == 1   # one row per ITEM, not per supplier line
    assert rows[0].prices["Harbour"].unit_price_ex_gst == D("65.20")
    assert rows[0].prices["Kiwi Frame"].unit_price_ex_gst == D("68.00")


def test_a_supplier_who_did_not_quote_an_item_is_simply_absent():
    harbour = _quote("Harbour", {"GIB-STD-10": (D("65.20"), D("652.00"))})
    kiwi = _quote("Kiwi Frame", {})   # didn't quote GIB at all

    rows = build_comparison_rows([GIB], [harbour, kiwi])

    assert "Kiwi Frame" not in rows[0].prices   # not a KeyError, not a None entry -- just absent


def test_out_of_rfq_lines_never_produce_a_comparison_row():
    supplier = _quote("Harbour", {"DELIVERY": (D("185.00"), D("185.00"))})   # not one of our items

    rows = build_comparison_rows([GIB], [supplier])

    assert rows[0].prices == {}   # DELIVERY isn't GIB-STD-10, so it never matches this row


# --- cheapest_supplier ---------------------------------------------------

def test_cheapest_supplier_by_unit_price():
    prices = {"Harbour": SupplierPrice(D("65.20"), None), "Kiwi Frame": SupplierPrice(D("63.50"), None)}
    assert cheapest_supplier(prices) == "Kiwi Frame"


def test_cheapest_supplier_ties_break_alphabetically_for_determinism():
    prices = {"Zeta": SupplierPrice(D("10.00"), None), "Alpha": SupplierPrice(D("10.00"), None)}
    assert cheapest_supplier(prices) == "Alpha"


def test_cheapest_supplier_of_an_unquoted_item_is_none():
    assert cheapest_supplier({}) is None


# --- supplier_grand_totals ---------------------------------------------------

def test_grand_totals_sum_only_what_each_supplier_actually_priced():
    rows = [
        ComparisonRow(GIB, {"Harbour": SupplierPrice(D("65.20"), D("652.00")),
                            "Kiwi Frame": SupplierPrice(D("68.00"), D("680.00"))}),
        ComparisonRow(MESH, {"Harbour": SupplierPrice(D("80.01"), D("1120.14"))}),   # Kiwi skipped this one
    ]
    totals = supplier_grand_totals(rows, ["Harbour", "Kiwi Frame"])
    assert totals["Harbour"] == D("1772.14")
    assert totals["Kiwi Frame"] == D("680.00")   # not penalised for the item it didn't quote


# --- export_workbook (integration: write, then read back) -------------------

def test_export_workbook_writes_a_demo_ready_excel(tmp_path: Path):
    harbour = _quote("Harbour", {"GIB-STD-10": (D("65.20"), D("652.00"))})
    kiwi = _quote("Kiwi Frame", {"GIB-STD-10": (D("68.00"), D("680.00"))})
    issues = [
        Issue("calc_error", Severity.ERROR, "14 x 80.01 != 11201.40", "Harbour", "MESH-665"),
        Issue("gst_unstated", Severity.WARNING, "no GST statement", "Kiwi Frame"),
    ]

    out = export_workbook([GIB], [harbour, kiwi], issues, tmp_path / "comparison.xlsx")

    wb = load_workbook(out)
    assert wb.sheetnames == ["Comparison", "Issues"]

    comp = wb["Comparison"]
    assert comp["A1"].value == "item_code"
    assert comp["D1"].value == "Harbour unit price"    # dynamic per-supplier columns
    assert comp["D2"].value == 65.20   # openpyxl reads Decimal cells back as float -- see the
                                        # GOTCHA comment in write_comparison_sheet
    assert comp["D2"].fill.fgColor.rgb[2:] == FILL_CHEAPEST.fgColor.rgb[2:]  # Harbour is cheaper
    assert comp["F2"].fill.fgColor.rgb == "00000000"     # Kiwi's cell is NOT highlighted

    issues_ws = wb["Issues"]
    assert issues_ws["A2"].value == "error"              # errors sorted before warnings
    assert issues_ws["A2"].fill.fgColor.rgb[2:] == FILL_ERROR.fgColor.rgb[2:]
    assert issues_ws["A3"].value == "warning"
    assert issues_ws["A3"].fill.fgColor.rgb[2:] == FILL_WARNING.fgColor.rgb[2:]


def test_missing_items_are_grouped_into_one_row_per_supplier():
    issues = [
        Issue("gst_unstated", Severity.WARNING, "no GST statement", "Kiwi Frame"),
        Issue("missing_item", Severity.WARNING, "not quoted", "Kiwi Frame", "MESH-665"),
        Issue("missing_item", Severity.WARNING, "not quoted", "Kiwi Frame", "GIB-STD-10"),
        Issue("missing_item", Severity.WARNING, "not quoted", "Northshore", "SEAL-123"),
    ]
    grouped = group_missing_items(issues)

    assert len(grouped) == 3
    kiwi = next(i for i in grouped if i.rule == "missing_item" and i.supplier == "Kiwi Frame")
    assert kiwi.item_code == "2 items"
    assert kiwi.message == "not quoted: GIB-STD-10, MESH-665"
    # A supplier missing only one item keeps its ordinary, specific row.
    assert Issue("missing_item", Severity.WARNING, "not quoted", "Northshore", "SEAL-123") in grouped
