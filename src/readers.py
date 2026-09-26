"""Turn each supplier file into a list of (line_no, text) pairs. Nothing more.

Rule for this whole module: carry the content over FAITHFULLY, do not interpret it.
No splitting "518.4 m2" into number and unit, no expanding "FR" to "framing". That is the
AI's job (extract.py) and the conversion code's job (normalise.py). If a reader "helpfully"
cleaned something and got it wrong, the mistake would be invisible: the original is gone.
"""
from dataclasses import dataclass
from pathlib import Path

import openpyxl
import pdfplumber


# frozen=True makes instances read-only: source lines are facts and must not be edited
# by later stages.
@dataclass(frozen=True)
class SourceLine:
    # Position of the line INSIDE the original file. This becomes quote_lines_raw.line_no,
    # the stable half of the upsert key from Phase 1. For Excel it is the sheet row number,
    # so it matches what you see when you open the file.
    line_no: int
    text: str


def read_excel(path: Path) -> list[SourceLine]:
    # pandas.read_excel would treat row 1 (the company name) as the header and misalign
    # everything below it. openpyxl gives raw cells, so we decide nothing about structure.
    # data_only=True returns the value Excel last calculated rather than the formula text
    # ("=C8*E8"), because the totals are exactly what later rules need to check.
    # Caveat: a workbook produced by a script and never opened in Excel has no cached
    # values, so formula cells would come back empty.
    wb = openpyxl.load_workbook(path, data_only=True)

    # Fail loudly instead of silently reading only the first sheet and dropping the rest.
    if len(wb.sheetnames) > 1:
        raise NotImplementedError(
            f"{path.name} has {len(wb.sheetnames)} sheets {wb.sheetnames}; "
            "only single-sheet quotes are supported")

    lines = []
    for row in wb.active.iter_rows():
        # Keep empty cells as "" so each value stays in its own column. If we dropped them,
        # a missing Qty would make Unit Price look like the Qty.
        cells = ["" if c.value is None else str(c.value).strip() for c in row]
        # Trailing empties add nothing; fully empty rows (spacers) are skipped, but the
        # row number is still the real one because we take it from the cell.
        while cells and cells[-1] == "":
            cells.pop()
        if cells:
            lines.append(SourceLine(row[0].row, " | ".join(cells)))
    return lines


def read_pdf(path: Path) -> list[SourceLine]:
    lines = []
    line_no = 0
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            # extract_text() returns None for a page with no text layer (e.g. a scan).
            text = page.extract_text() or ""
            for raw in text.splitlines():
                # Count blank lines too so numbering reflects the true position, and keeps
                # the same numbers if a blank line is later treated differently.
                line_no += 1
                if raw.strip():
                    lines.append(SourceLine(line_no, raw.strip()))

    # A scanned PDF is just a picture. Returning an empty list would look like "supplier
    # quoted nothing", which is a wrong business conclusion, so stop here instead.
    if not lines:
        raise ValueError(f"{path.name} has no text layer (scanned image?); needs OCR or a vision model")
    return lines


def read_text(path: Path) -> list[SourceLine]:
    # Strict utf-8 (no errors="ignore"): garbled input should fail here, not turn into
    # silently wrong prices further down.
    text = path.read_text(encoding="utf-8")
    return [SourceLine(n, raw.strip())
            for n, raw in enumerate(text.splitlines(), start=1) if raw.strip()]


# Dispatch table: adding a new file type means adding one entry, not editing an if/elif chain.
READERS = {
    ".xlsx": read_excel,
    ".pdf": read_pdf,
    ".txt": read_text,
}


def read_quote(path: Path) -> list[SourceLine]:
    path = Path(path)
    reader = READERS.get(path.suffix.lower())
    if reader is None:
        raise ValueError(f"Unsupported file type {path.suffix!r} for {path.name}; "
                         f"supported: {sorted(READERS)}")
    return reader(path)
