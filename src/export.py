"""FR-7: export the comparison Excel. Pure data-shaping, then a thin openpyxl-writing layer.

Like rules.py, this module never opens a database connection or calls the AI: it receives
already-assembled `SupplierQuote`/`SupplierLine` objects (the same shapes rules.py uses) and
an `Issue` list (rules.detect_issues()'s own output). Phase 8's pipeline.py is the only place
that queries the database and wires these pieces together -- keeping that wiring in one spot
means export.py can be tested the same cheap way rules.py is: build objects in memory, no I/O.
"""
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from src.rules import Issue, Severity, SupplierQuote
from src.schemas import OUT_OF_RFQ


@dataclass(frozen=True)
class Item:
    """One RFQ row's display info. Just enough to label the comparison sheet -- rfq_qty and
    ref_price_ex_gst live in the `items` table too, but rules.py already owns the checks that
    need them (qty_short, price_band's fallback baseline), so this module doesn't repeat them.
    """
    item_code: str
    description: str
    unit: str


@dataclass(frozen=True)
class SupplierPrice:
    unit_price_ex_gst: Decimal
    line_total_ex_gst: Decimal | None   # None for a matched-but-priceless line (rare, but real)


@dataclass(frozen=True)
class ComparisonRow:
    item: Item
    # Keyed by supplier name. A supplier absent from this dict simply didn't quote the item --
    # rules.py's missing_item rule already reports that on the Issues sheet, so this sheet just
    # renders a blank cell rather than repeating the finding.
    prices: dict[str, SupplierPrice]


# ---------------------------------------------------------------------------------
# Pure data shaping: long (one row per supplier line) -> wide (one row per item)
# ---------------------------------------------------------------------------------

def build_comparison_rows(items: list[Item], quotes: list[SupplierQuote]) -> list[ComparisonRow]:
    """The "pivot": SQL can't declare one column per supplier ahead of time because the
    supplier count isn't fixed, so the reshape happens here in Python instead, where a dict
    can grow one key per supplier with no schema to change.
    """
    rows = []
    for item in items:
        prices: dict[str, SupplierPrice] = {}
        for quote in quotes:
            for line in quote.lines:
                # OUT_OF_RFQ and unmatched (None) lines don't belong to any RFQ row; they
                # only ever appear on the Issues sheet (out_of_rfq / lump_sum), never here.
                if line.item_code != item.item_code:
                    continue
                if line.unit_price_ex_gst is None:
                    continue   # matched but no usable price (e.g. extraction gap) -> leave blank
                # If a supplier somehow has two lines for the same item, keep the first and
                # ignore the rest: the mock data never does this, and silently summing them
                # would be guessing at a business rule (bulk discount? duplicate entry?) that
                # isn't ours to invent. A real second occurrence should have failed earlier,
                # at the (quote_id, line_no) -> item_code matching step, not been patched here.
                if quote.supplier not in prices:
                    prices[quote.supplier] = SupplierPrice(
                        line.unit_price_ex_gst, line.line_total_ex_gst)
        rows.append(ComparisonRow(item, prices))
    return rows


def cheapest_supplier(prices: dict[str, SupplierPrice]) -> str | None:
    """None if nobody quoted this item. Ties broken alphabetically by supplier name so the
    result is deterministic -- min() over floats/Decimals with equal values would otherwise
    silently depend on dict ordering, which is an easy source of a flaky-looking demo.
    """
    if not prices:
        return None
    return min(prices, key=lambda supplier: (prices[supplier].unit_price_ex_gst, supplier))


def supplier_grand_totals(rows: list[ComparisonRow], suppliers: list[str]) -> dict[str, Decimal]:
    """Each supplier's total across every item they priced -- the "grand total" column. Items a
    supplier didn't quote just contribute nothing, same as a blank cell contributes nothing
    to a SUM() in Excel.
    """
    totals = {supplier: Decimal(0) for supplier in suppliers}
    for row in rows:
        for supplier, price in row.prices.items():
            if price.line_total_ex_gst is not None:
                totals[supplier] += price.line_total_ex_gst
    return totals


def group_missing_items(issues: list[Issue]) -> list[Issue]:
    """Collapse each supplier's missing_item findings into ONE row for display.

    rules.py reports one finding per item on purpose: that's the data, and evaluate.py scores
    it item by item. But for a reader, "Kiwi Frame doesn't sell plasterboard, insulation, ..."
    is one fact, not 11 things to act on -- 11 near-identical yellow rows bury the 4 warnings
    that actually need a decision. Presentation is this module's job, so the grouping lives
    here. A supplier missing just one item keeps its normal row.
    """
    missing_by_supplier: dict[str, list[str]] = {}
    for issue in issues:
        if issue.rule == "missing_item":
            missing_by_supplier.setdefault(issue.supplier, []).append(issue.item_code or "?")

    grouped: list[Issue] = []
    emitted: set[str] = set()
    for issue in issues:
        if issue.rule != "missing_item":
            grouped.append(issue)
            continue
        codes = missing_by_supplier[issue.supplier]
        if len(codes) == 1:
            grouped.append(issue)
        elif issue.supplier not in emitted:   # first one stands in for the whole group
            emitted.add(issue.supplier)
            grouped.append(Issue("missing_item", issue.severity,
                                 f"not quoted: {', '.join(sorted(codes))}",
                                 issue.supplier, f"{len(codes)} items"))
    return grouped


# ---------------------------------------------------------------------------------
# openpyxl writing: the only part of this module with I/O
# ---------------------------------------------------------------------------------

# Same red/yellow/green vocabulary Excel's own conditional formatting uses, so a reader who
# has never seen this tool still knows red = bad on sight.
FILL_CHEAPEST = PatternFill("solid", fgColor="C6EFCE")   # Excel's built-in "Good" green
FILL_ERROR = PatternFill("solid", fgColor="FFC7CE")      # Excel's built-in "Bad" red
FILL_WARNING = PatternFill("solid", fgColor="FFEB9C")    # Excel's built-in "Neutral" yellow
FILL_INFO = PatternFill("solid", fgColor="DDEBF7")       # a plain, calm blue for FYI rows
HEADER_FONT = Font(bold=True)

SEVERITY_FILL = {Severity.ERROR: FILL_ERROR, Severity.WARNING: FILL_WARNING, Severity.INFO: FILL_INFO}


def _autofit(ws: Worksheet, n_columns: int) -> None:
    # GOTCHA: openpyxl never measures text and never auto-sizes columns -- Excel does that
    # live in the app, but a file written by code has no app running, so every column comes
    # out at the default width unless you set one yourself. This is the cheapest fix: widest
    # cell's character count plus a little padding, not pixel-perfect but good enough to demo.
    for col_idx in range(1, n_columns + 1):
        letter = ws.cell(row=1, column=col_idx).column_letter
        widest = max((len(str(ws.cell(row=r, column=col_idx).value or "")) for r in range(1, ws.max_row + 1)),
                     default=0)
        ws.column_dimensions[letter].width = min(widest + 2, 40)


def write_comparison_sheet(ws: Worksheet, rows: list[ComparisonRow], suppliers: list[str]) -> None:
    # GOTCHA: openpyxl silently converts every Decimal cell value to float on write -- XLSX's
    # numeric type has no concept of Decimal, so there is no error, just a quiet downgrade.
    # Safe here only because Phase 3 already rounded these to cents before they got this far;
    # this module writes a display copy, it never feeds the float back into more arithmetic.
    headers = ["item_code", "description", "unit"]
    for supplier in suppliers:
        headers += [f"{supplier} unit price", f"{supplier} subtotal"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT
    ws.freeze_panes = "A2"   # header row stays visible while scrolling through 20 items

    for row in rows:
        cheapest = cheapest_supplier(row.prices)
        out = [row.item.item_code, row.item.description, row.item.unit]
        for supplier in suppliers:
            price = row.prices.get(supplier)
            out += [price.unit_price_ex_gst if price else None,
                   price.line_total_ex_gst if price else None]
        ws.append(out)
        # Column 4 is the first supplier's unit-price column; each supplier occupies 2
        # columns (price, subtotal), so supplier i's price column is 4 + 2*i.
        if cheapest is not None:
            col = 4 + 2 * suppliers.index(cheapest)
            ws.cell(row=ws.max_row, column=col).fill = FILL_CHEAPEST

    totals = supplier_grand_totals(rows, suppliers)
    total_row = ["TOTAL", "", ""]
    for supplier in suppliers:
        total_row += [None, totals[supplier]]   # blank under "unit price": a total has no unit price
    ws.append(total_row)
    for cell in ws[ws.max_row]:
        cell.font = HEADER_FONT

    _autofit(ws, len(headers))


def write_issues_sheet(ws: Worksheet, issues: list[Issue]) -> None:
    headers = ["severity", "rule", "supplier", "item_code", "message"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT
    ws.freeze_panes = "A2"

    # Errors first: a reader scanning top-down should hit the must-fix findings before the
    # for-your-information ones. Severity.ERROR/WARNING/INFO are declared in that order in
    # rules.py's Enum, so .value sorting isn't reliable (alphabetical: error, info, warning) --
    # sort by the enum member's declaration order instead.
    order = {Severity.ERROR: 0, Severity.WARNING: 1, Severity.INFO: 2}
    for issue in sorted(group_missing_items(issues), key=lambda i: order[i.severity]):
        ws.append([issue.severity.value, issue.rule, issue.supplier, issue.item_code, issue.message])
        for cell in ws[ws.max_row]:
            cell.fill = SEVERITY_FILL[issue.severity]

    _autofit(ws, len(headers))


def export_workbook(items: list[Item], quotes: list[SupplierQuote], issues: list[Issue],
                    output_path: Path) -> Path:
    """Build both sheets and save. The one function Phase 8's pipeline.py actually calls."""
    suppliers = [quote.supplier for quote in quotes]
    rows = build_comparison_rows(items, quotes)

    wb = Workbook()
    comparison_ws = wb.active
    comparison_ws.title = "Comparison"
    write_comparison_sheet(comparison_ws, rows, suppliers)

    issues_ws = wb.create_sheet("Issues")
    write_issues_sheet(issues_ws, issues)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path
